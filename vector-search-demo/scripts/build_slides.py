"""Generates the companion slide deck for the Vector Search demo video.

Run: python3 scripts/build_slides.py
Output: ../VectorSearch_Semantic_Indexing.pptx (one level up from scripts/)
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

OUTPUT_PATH = os.path.join(os.path.dirname(__file__), "..", "VectorSearch_Semantic_Indexing.pptx")

NAVY = RGBColor(0x1F, 0x2A, 0x44)
BLUE = RGBColor(0x2B, 0x5F, 0xD9)
GRAY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def add_title(slide, text, size=36, color=NAVY, top=Inches(0.4)):
    box = slide.shapes.add_textbox(Inches(0.6), top, Inches(12.1), Inches(1.0))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = True
    p.font.color.rgb = color
    return box


def add_bullets(slide, items, top=Inches(1.5), left=Inches(0.8), width=Inches(11.7),
                 height=Inches(5.3), size=22, color=GRAY):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p.text = ("•  " if level == 0 else "–  ") + text
        p.level = level
        p.font.size = Pt(size if level == 0 else size - 2)
        p.font.color.rgb = color
        p.space_after = Pt(10)
    return box


def set_background(slide, color):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def add_image_slide(title_text, caption_items, image_path):
    slide = add_slide()
    add_title(slide, title_text)
    add_bullets(slide, caption_items, top=Inches(1.4), height=Inches(1.6), size=18)
    slide.shapes.add_picture(image_path, Inches(1.9), Inches(2.9), height=Inches(4.3))
    return slide


# ---------------------------------------------------------------- Slide 1
s = add_slide()
set_background(s, NAVY)
box = s.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11.7), Inches(1.5))
p = box.text_frame.paragraphs[0]
p.text = "Vector Search & Semantic Indexing in Full-Stack Apps"
p.font.size = Pt(40)
p.font.bold = True
p.font.color.rgb = WHITE
p.alignment = PP_ALIGN.LEFT

box2 = s.shapes.add_textbox(Inches(0.8), Inches(4.1), Inches(11.7), Inches(1.0))
p2 = box2.text_frame.paragraphs[0]
p2.text = "Querying data by meaning, not just by matching strings"
p2.font.size = Pt(22)
p2.font.color.rgb = RGBColor(0xC9, 0xD4, 0xEE)

box3 = s.shapes.add_textbox(Inches(0.8), Inches(6.6), Inches(11.7), Inches(0.6))
p3 = box3.text_frame.paragraphs[0]
p3.text = "LIS 467 — Client/Server & Full-Stack Architectures"
p3.font.size = Pt(16)
p3.font.color.rgb = RGBColor(0x8A, 0x97, 0xBC)

# ---------------------------------------------------------------- Slide 2: The Concept
s = add_slide()
add_title(s, "The Concept")
add_bullets(s, [
    "Traditional queries (SQL WHERE, Mongo $regex) match exact strings or structured fields.",
    "Modern AI-driven apps query by semantic similarity instead.",
    "Text, images, or documents are converted into high-dimensional numeric vectors called embeddings.",
    "Vectors that are close together in that space represent similar meaning — even with zero shared words.",
    ("\"organizing books so people can find them\" \u2248 \"cataloging and classification in libraries\"", 1),
    "Similarity is measured with cosine similarity: the cosine of the angle between two vectors.",
])

# ---------------------------------------------------------------- Slide 3: Why it fits
s = add_slide()
add_title(s, "Why This Fits Full-Stack Development")
add_bullets(s, [
    "Deepens backend database and document-store skills beyond CRUD.",
    "Real systems: MongoDB Atlas Vector Search, PostgreSQL + pgvector, Pinecone, Milvus.",
    "Pairs naturally with Node.js / Express backends you already know how to build.",
    "This is the retrieval layer behind Retrieval-Augmented Generation (RAG) — how LLM apps",
    ("ground answers in your own data instead of relying only on training data.", 1),
    "Same skills apply across the stack: JSON documents in \u2192 embeddings \u2192 vector index \u2192 API endpoint.",
])

# ---------------------------------------------------------------- Slide 4: Architecture
s = add_slide()
add_title(s, "Architecture: How It Works")
add_bullets(s, [
    "1. Indexing (offline / build step)",
    ("JSON documents \u2192 embedding model \u2192 384-dim vectors \u2192 stored alongside each document", 1),
    "2. Query time (API request)",
    ("User query \u2192 same embedding model \u2192 query vector", 1),
    ("Compare query vector to every stored vector using cosine similarity", 1),
    ("Sort by score, return top-K most similar documents", 1),
    "3. (Production) Approximate Nearest Neighbor index (HNSW) replaces brute-force scan at scale",
])

# ---------------------------------------------------------------- Slide 4b: MATLAB visualization
MATLAB_PLOT = os.path.join(os.path.dirname(__file__), "..", "matlab", "vector_space_plot.png")
add_image_slide(
    "Visualizing the Vector Space (MATLAB)",
    [
        "Companion MATLAB script builds the same idea from scratch: TF-IDF vectors instead of a neural model.",
        "SVD (the math behind PCA) reduces the high-dimensional space to 2D so it can be plotted.",
        "Documents from the same category cluster together — visual proof that distance in vector space reflects meaning.",
    ],
    MATLAB_PLOT,
)

# ---------------------------------------------------------------- Slide 5: Exercise
s = add_slide()
add_title(s, "Class Exercise")
add_bullets(s, [
    "Goal: implement a retrieval-augmented search endpoint in Node.js.",
    "Dataset: small collection of JSON documents (title, category, text).",
    "Step 1 — Embed: convert each document's text into a vector using a local embedding model",
    ("(transformers.js, no API key needed).", 1),
    "Step 2 — Index: store {document, embedding} pairs (data/index.json).",
    "Step 3 — Retrieve: embed the incoming query, compute cosine similarity against every",
    ("stored vector, return the top-K ranked results.", 1),
    "Step 4 — Compare: run the same query through naive keyword matching and observe the gap.",
])

# ---------------------------------------------------------------- Slide 6: Demo transition
s = add_slide()
set_background(s, BLUE)
box = s.shapes.add_textbox(Inches(1.0), Inches(3.0), Inches(11.3), Inches(1.5))
p = box.text_frame.paragraphs[0]
p.text = "Live Demo"
p.font.size = Pt(44)
p.font.bold = True
p.font.color.rgb = WHITE
box2 = s.shapes.add_textbox(Inches(1.0), Inches(4.2), Inches(11.3), Inches(1.0))
p2 = box2.text_frame.paragraphs[0]
p2.text = "Node.js + Express + local embeddings, semantic vs. keyword search side-by-side"
p2.font.size = Pt(20)
p2.font.color.rgb = RGBColor(0xDD, 0xE6, 0xFC)

# ---------------------------------------------------------------- Slide 7: Key takeaways
s = add_slide()
add_title(s, "Key Takeaways")
add_bullets(s, [
    "Semantic search retrieves by meaning; keyword search retrieves by exact substring.",
    "Embeddings + cosine similarity are the core primitives behind modern AI search and RAG.",
    "These concepts map directly onto production tools: pgvector, MongoDB Atlas Vector Search, Pinecone.",
    "A brute-force cosine scan works for small datasets; production systems need an ANN index (HNSW).",
    "Full-stack takeaway: the backend pattern is the same you already know —",
    ("JSON in, API endpoint out — only the storage and matching logic changes.", 1),
])

prs.save(OUTPUT_PATH)
print(f"Saved {os.path.abspath(OUTPUT_PATH)}")
